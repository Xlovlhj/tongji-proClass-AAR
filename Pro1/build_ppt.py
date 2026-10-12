"""Build an editable, ten-slide raa-v0 presentation from completed runs.

Run from the project root: python3 Pro1/build_ppt.py
No simulation is launched and no experiment outputs are overwritten.
Text, tables, flowchart, outline plots and time-series chart are native PPTX
objects. Actual Isaac screenshots are raster pictures; eight original MP4
videos are embedded in the PPTX, each with an editable explanation below it.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "ppt"
W, H = 13.333333, 7.5
TOTAL_SLIDES = 10
FONT = "Noto Sans CJK SC"
BG = "FAFBFC"
INK = "20343D"
MUTED = "65747B"
TEAL = "008C82"
PALE = "EAF5F3"
GRID = "DCE4E7"
TARGET = "8D9AA2"
ORANGE = "C77936"
BLUE = "5075AA"
WHITE = "FFFFFF"
CONTROL = ("pd", "gc-pd", "rbf", "fuzzy")
LABEL = {"pd": "PD", "gc-pd": "重力补偿 PD", "rbf": "RBF 自适应", "fuzzy": "模糊自适应"}
SHORT = {"pd": "PD", "gc-pd": "重力补偿 PD", "rbf": "RBF", "fuzzy": "模糊"}
SERIES_COLOR = {"pd": INK, "gc-pd": BLUE, "rbf": ORANGE, "fuzzy": TEAL}
SERIES_DASH = {"pd": MSO_LINE_DASH_STYLE.SOLID, "gc-pd": MSO_LINE_DASH_STYLE.DASH,
               "rbf": MSO_LINE_DASH_STYLE.LONG_DASH, "fuzzy": MSO_LINE_DASH_STYLE.SOLID}
PPT_NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
VIDEO_SUBTITLE_OLD = "点击视频画面播放；每个视频对应一种控制器。"
VIDEO_SUBTITLE = "进入本页后，四个视频自动同时播放。"


def rgb(value):
    return RGBColor.from_string(value)


def set_font(run, size, color=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    props = run._r.get_or_add_rPr()
    props.set("lang", "zh-CN")
    for name in ("a:ea", "a:cs"):
        node = props.find(name, props.nsmap)
        if node is None:
            node = OxmlElement(name)
            props.append(node)
        node.set("typeface", FONT)


def text(slide, x, y, w, h, content, size=22, color=INK, bold=False,
         align=PP_ALIGN.LEFT, vertical=MSO_ANCHOR.TOP, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or content.split("\n")[0][:48]
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.015)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = vertical
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(5)
        p.line_spacing = 1.12
        run = p.add_run()
        run.text = line
        set_font(run, size, color, bold)
    return shape


def rect(slide, x, y, w, h, fill=WHITE, line_color=None, radius=False, name=""):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0] = 0.05
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line_color:
        s.line.color.rgb = rgb(line_color)
        s.line.width = Pt(0.7)
    else:
        s.line.fill.background()
    if name:
        s.name = name
    return s


def line(slide, x1, y1, x2, y2, color=GRID, width=1, dash=None):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if dash:
        s.line.dash_style = dash
    return s


def picture(slide, source, x, y, w, h):
    # These project viewport captures have a 4:3 aspect ratio; preserve it.
    return slide.shapes.add_picture(str(source), Inches(x), Inches(y),
                                    width=Inches(w), height=Inches(h))


def base(prs, number, title=None, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0.6, 0.53, 0.075, 0.43, TEAL)
    if title:
        text(s, 0.82, 0.48, 11.82, 0.64, title, 30, bold=True)
    if subtitle:
        text(s, 0.82, 1.13, 11.7, 0.45, subtitle, 18, MUTED)
    line(s, 0.6, 7.02, 12.73, 7.02, GRID, 0.6)
    text(s, 0.6, 7.12, 8.2, 0.2, "robot arm artist  ·  raa-v0", 9, MUTED)
    text(s, 12.0, 7.10, 0.73, 0.25, f"{number:02d} / {TOTAL_SLIDES:02d}", 10, MUTED, align=PP_ALIGN.RIGHT)
    return s


def table(slide, x, y, w, h, rows, widths=None, size=20, highlight_row=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                  Inches(w), Inches(h))
    tab = shape.table
    if widths:
        for col, fraction in zip(tab.columns, widths):
            col.width = Inches(w * fraction)
    for r, values in enumerate(rows):
        tab.rows[r].height = Inches(h / len(rows))
        for c, value in enumerate(values):
            cell = tab.cell(r, c)
            cell.text = str(value)
            cell.margin_left = Inches(0.13)
            cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.055)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PALE if r == 0 or r == highlight_row else WHITE)
            # Explicit cell borders avoid template-dependent decoration.
            tcpr = cell._tc.get_or_add_tcPr()
            for tag in ("lnL", "lnR", "lnT", "lnB"):
                node = OxmlElement("a:" + tag)
                node.set("w", "6350")
                solid = OxmlElement("a:solidFill")
                srgb = OxmlElement("a:srgbClr")
                srgb.set("val", GRID)
                solid.append(srgb)
                node.append(solid)
                tcpr.append(node)
            for p in cell.text_frame.paragraphs:
                p.space_before = p.space_after = Pt(0)
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    set_font(run, size, TEAL if r == highlight_row else INK, r == 0)
    return shape


def rdp(points, epsilon=0.05):
    """Simplify only the displayed mm polyline, never the numeric metrics."""
    if len(points) < 3:
        return points
    a, b = points[0], points[-1]
    dx, dy = b[0] - a[0], b[1] - a[1]
    denom = dx * dx + dy * dy
    largest, index = -1.0, 0
    for i in range(1, len(points) - 1):
        p = points[i]
        t = max(0.0, min(1.0, ((p[0] - a[0]) * dx + (p[1] - a[1]) * dy) / denom)) if denom else 0
        distance = math.hypot(p[0] - a[0] - t * dx, p[1] - a[1] - t * dy)
        if distance > largest:
            largest, index = distance, i
    if largest <= epsilon:
        return [a, b]
    return rdp(points[:index + 1], epsilon)[:-1] + rdp(points[index:], epsilon)


def read_strokes(path):
    target, actual = {}, {}
    current, previous = None, None
    with path.open(newline="") as stream:
        reader = csv.reader(stream)
        fields = next(reader)
        at = {name: fields.index(name) for name in
              ("target_x", "target_y", "actual_x", "actual_y", "pen_down", "stroke_index")}
        for index, row in enumerate(reader):
            down = row[at["pen_down"]] in ("True", "true", "1")
            stroke = int(row[at["stroke_index"]]) if down else None
            if stroke != current:
                if current is not None and previous:
                    target[current].append(previous[0])
                    actual[current].append(previous[1])
                current = stroke
                previous = None
            if not down:
                continue
            point = ((float(row[at["target_x"]]) * 1000, float(row[at["target_y"]]) * 1000),
                     (float(row[at["actual_x"]]) * 1000, float(row[at["actual_y"]]) * 1000))
            if stroke not in target:
                target[stroke], actual[stroke] = [], []
            if not target[stroke] or index % 10 == 0:
                target[stroke].append(point[0])
                actual[stroke].append(point[1])
            previous = point
    if current is not None and previous:
        target[current].append(previous[0])
        actual[current].append(previous[1])
    assert len(target) == len(actual) == 8, path
    return {"target": [rdp(target[k]) for k in sorted(target)],
            "actual": [rdp(actual[k]) for k in sorted(actual)]}


def polyline(slide, points, color=TEAL, width=1.4, dash=None, name="editable trajectory"):
    if len(points) < 2:
        return None
    # Local integer coordinates avoid rounding paths in the inch coordinate system.
    sx, sy = points[0]
    builder = slide.shapes.build_freeform(round(sx * 10000), round(sy * 10000),
                                          scale=Inches(1) / 10000)
    builder.add_line_segments([(round(x * 10000), round(y * 10000)) for x, y in points[1:]], close=False)
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    shape.name = name
    return shape


def legend(slide, x, y, entries, size=16, spacing=1.26):
    for i, (label, color, dash) in enumerate(entries):
        offset = x + i * spacing
        line(slide, offset, y + 0.13, offset + 0.34, y + 0.13, color, 1.8, dash)
        text(slide, offset + 0.42, y - 0.035, spacing - 0.42, 0.35, label, size, MUTED)


def outline_plot(slide, x, y, w, h, title, entries, *, small=False, axes=True):
    # Every trajectory chart has the same mm domain and equal spatial aspect.
    title_h = 0.40
    text(slide, x, y, w, title_h, title, 20 if small else 22, bold=True,
         align=PP_ALIGN.CENTER)
    if axes:
        reserve_left, reserve_bottom = (0.53, 0.61) if small else (0.64, 0.63)
        side = min(w - reserve_left - 0.16, h - title_h - reserve_bottom)
        px = x + reserve_left + (w - reserve_left - 0.16 - side) / 2
        py = y + title_h + 0.08
        xmin, xmax, ymin, ymax = 320.0, 580.0, -130.0, 130.0
        ticks_x, ticks_y = (350, 450, 550), (-100, 0, 100)
        label_size = 13 if small else 16
        for value in ticks_x:
            xx = px + (value - xmin) / (xmax - xmin) * side
            line(slide, xx, py, xx, py + side, GRID, 0.45)
            text(slide, xx - 0.25, py + side + 0.03, 0.5, 0.27, str(value), label_size,
                 MUTED, align=PP_ALIGN.CENTER)
        for value in ticks_y:
            yy = py + (ymax - value) / (ymax - ymin) * side
            line(slide, px, yy, px + side, yy, GRID, 0.45)
            text(slide, px - 0.70, yy - 0.135, 0.60, 0.28, str(value), label_size,
                 MUTED, align=PP_ALIGN.RIGHT)
        line(slide, px, py, px, py + side, TARGET, 0.8)
        line(slide, px, py + side, px + side, py + side, TARGET, 0.8)
        text(slide, px, py + side + 0.32, side, 0.30, "X / mm", label_size + 1,
             MUTED, align=PP_ALIGN.CENTER)
        text(slide, px - 0.57, py - 0.13, 0.90, 0.30, "Y / mm", label_size, MUTED)
    else:
        side = min(w, h - title_h - 0.12)
        px, py = x + (w - side) / 2, y + title_h + 0.12
        xmin, xmax, ymin, ymax = 360.0, 540.0, -90.0, 90.0
    for label, strokes, color, dash in entries:
        for index, stroke in enumerate(strokes):
            for point in stroke:
                assert xmin <= point[0] <= xmax and ymin <= point[1] <= ymax, (label, point)
            points = [(px + (a - xmin) / (xmax - xmin) * side,
                       py + (ymax - b) / (ymax - ymin) * side) for a, b in stroke]
            polyline(slide, points, color, 1.35 if small else 1.7, dash,
                     f"{title} / {label} / stroke {index + 1}")


def read_deviation(path):
    values = []
    with path.open(newline="") as stream:
        reader = csv.reader(stream)
        fields = next(reader)
        ti, di = fields.index("time"), fields.index("disturbance_deviation_m")
        for index, row in enumerate(reader):
            t = float(row[ti])
            if t > 27:
                break
            if t >= 19 and index % 10 == 0:
                values.append((t, float(row[di]) * 1000))
    assert len(values) == 801
    return values


def deviation_chart(slide, x, y, w, h, deviations):
    text(slide, x, y - 0.48, w, 0.4, "第一次推力前后的额外偏移", 21, bold=True)
    data = XyChartData()
    for controller in CONTROL:
        series = data.add_series(SHORT[controller])
        for t, v in deviations[controller]:
            series.add_data_point(t, v)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
                                  Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_title = False
    chart.chart_style = 2
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name = FONT
    chart.legend.font.size = Pt(14)
    chart.legend.font.color.rgb = rgb(MUTED)
    for controller, series in zip(CONTROL, chart.series):
        series.format.line.color.rgb = rgb(SERIES_COLOR[controller])
        series.format.line.width = Pt(1.6)
        series.format.line.dash_style = SERIES_DASH[controller]
        series.smooth = False
    xaxis, yaxis = chart.category_axis, chart.value_axis
    for axis in (xaxis, yaxis):
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(14)
        axis.tick_labels.font.color.rgb = rgb(MUTED)
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = rgb(TARGET)
        axis.format.line.width = Pt(0.7)
        axis.has_title = True
    xaxis.minimum_scale, xaxis.maximum_scale, xaxis.major_unit = 19, 27, 2
    maximum = max(value for points in deviations.values() for _, value in points)
    yaxis.minimum_scale, yaxis.maximum_scale, yaxis.major_unit = 0, math.ceil(maximum / 2) * 2, 2
    yaxis.has_major_gridlines = True
    yaxis.major_gridlines.format.line.color.rgb = rgb(GRID)
    yaxis.major_gridlines.format.line.width = Pt(0.45)
    for axis, content in ((xaxis, "时间（s）"), (yaxis, "三维位置差（mm）")):
        axis.axis_title.text_frame.text = content
        for p in axis.axis_title.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, 15, MUTED)
    return chart


def load_notes():
    source = OUT / "speaker_notes.md"
    if not source.exists():
        return {}
    notes = {}
    parts = re.split(r"(?m)^##\s+", source.read_text())
    for part in parts[1:]:
        first, _, body = part.partition("\n")
        match = re.match(r"(?:第\s*)?(\d+)", first)
        if match:
            notes[int(match.group(1))] = body.strip()
    return notes


def configure_parallel_video_autoplay(slide_root):
    """Start all four media nodes together at slide-root time zero.

    Native p:video nodes are siblings in one parallel time container. Changing
    their start conditions from indefinite (manual) to zero starts them with
    the slide, rather than sequencing them through individual clicks. Keep
    media durations implicit so each MP4 uses its actual playback duration.
    """
    children = slide_root.find(
        'p:timing/p:tnLst/p:par/p:cTn[@nodeType="tmRoot"]/p:childTnLst', PPT_NS)
    if children is None:
        raise ValueError("Video slide has no shared parallel timing root")
    videos = children.findall("p:video", PPT_NS)
    if len(videos) != 4:
        raise ValueError("Each video slide must have four parallel video nodes")
    shape_ids = []
    for video in videos:
        ctn = video.find("p:cMediaNode/p:cTn", PPT_NS)
        conditions = ctn.find("p:stCondLst", PPT_NS) if ctn is not None else None
        target = video.find("p:cMediaNode/p:tgtEl/p:spTgt", PPT_NS)
        if conditions is None or target is None:
            raise ValueError("Video timing requires a start condition and shape target")
        starts = conditions.findall("p:cond", PPT_NS)
        if len(starts) != 1 or starts[0].get("evt") is not None or len(starts[0]):
            raise ValueError("Unexpected interactive start condition; cannot silently replace it")
        starts[0].set("delay", "0")
        shape_ids.append(target.get("spid"))
    if len(set(shape_ids)) != 4:
        raise ValueError("The four video nodes must target distinct movie shapes")
    return shape_ids


def synchronize_existing_presentation(destination):
    """Patch playback in the existing PPTX, preserving the user's other edits."""
    if not destination.is_file():
        raise FileNotFoundError(destination)
    temporary = destination.with_name(f".{destination.stem}.syncing.pptx")
    changed, targets, original_media_hashes = [], {}, {}
    text_replacements = {
        VIDEO_SUBTITLE_OLD: VIDEO_SUBTITLE,
        "放映时直接点击画面播放": "进入本页后，四个视频自动同时播放",
        "放映时点击各自的画面即可播放": "进入本页后四个视频自动同时播放",
        "点击画面即可播放": "进入本页后四个视频自动同时播放",
    }
    xml_names = {f"ppt/slides/slide{i}.xml" for i in (8, 9)}
    xml_names.update(f"ppt/notesSlides/notesSlide{i}.xml" for i in (8, 9))
    try:
        with ZipFile(destination) as source, ZipFile(temporary, "w") as result:
            for item in source.infolist():
                if item.filename in xml_names:
                    original = source.read(item.filename)
                    root = etree.fromstring(original)
                    if item.filename.startswith("ppt/slides/"):
                        targets[item.filename] = configure_parallel_video_autoplay(root)
                    for element in root.findall(".//a:t", PPT_NS):
                        if element.text:
                            for old, new in text_replacements.items():
                                element.text = element.text.replace(old, new)
                    updated = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    result.writestr(item, updated)
                    if original != updated:
                        changed.append(item.filename)
                else:
                    # Copy the complete existing media and other package parts.
                    # Video bytes are not decoded or re-encoded.
                    digest = hashlib.sha256() if item.filename.endswith(".mp4") else None
                    with source.open(item) as incoming, result.open(item, "w") as outgoing:
                        while chunk := incoming.read(1024 * 1024):
                            outgoing.write(chunk)
                            if digest is not None:
                                digest.update(chunk)
                    if digest is not None:
                        original_media_hashes[item.filename] = digest.hexdigest()
        if len(targets) != 2 or len(original_media_hashes) != 8:
            raise ValueError("Expected two video slides and eight embedded MP4 files")
        with ZipFile(temporary) as result:
            for name, expected in original_media_hashes.items():
                with result.open(name) as stream:
                    actual = hashlib.file_digest(stream, "sha256").hexdigest()
                if actual != expected:
                    raise ValueError(f"Embedded video changed unexpectedly: {name}")
            for name in targets:
                root = etree.fromstring(result.read(name))
                nodes = root.findall(".//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond", PPT_NS)
                if len(nodes) != 4 or any(n.get("delay") != "0" or n.get("evt") for n in nodes):
                    raise ValueError(f"Parallel zero-delay playback validation failed: {name}")
        temporary.replace(destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    report = {"presentation": str(destination), "video_slides": [8, 9],
              "mode": "four videos automatically start together on slide entry",
              "start_delay_ms": 0, "parallel_targets": targets,
              "modified_package_entries": changed,
              "embedded_media_sha256": original_media_hashes,
              "embedded_videos_unchanged": True,
              "pptx_sha256": hashlib.file_digest(destination.open("rb"), "sha256").hexdigest()}
    (OUT / "video_playback_verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    sources_file = OUT / "presentation_sources.json"
    if sources_file.exists():
        sources = json.loads(sources_file.read_text())
        sources["video_delivery"] = "8 complete internally embedded MP4 files; four videos start automatically in parallel on each video slide; no external video links"
        sources["video_playback"] = {"trigger": "slide entry", "start_delay_ms": 0,
                                      "mode": "parallel", "videos_per_slide": 4, "slides": [8, 9]}
        sources_file.write_text(json.dumps(sources, ensure_ascii=False, indent=2) + "\n")
    print("Updated both video slides: four parallel videos, zero start delay, all eight MP4s unchanged.")
    return report


def video_slide(prs, number, condition, source_paths):
    normal = condition == "normal"
    title = "正常绘画：四种控制器的视频" if normal else "外力干扰：四种控制器的视频"
    subtitle = VIDEO_SUBTITLE
    slide = base(prs, number, title, subtitle)
    explanations = {
        "normal": {
            "pd": "按位置偏差纠正。\n自重使轮廓偏离明显。",
            "gc-pd": "先抵消机械臂自重。\n轮廓更贴近目标线。",
            "rbf": "边画边学习补偿。\n轮廓接近目标线。",
            "fuzzy": "根据误差规则补偿。\n本次整体误差最小。",
        },
        "disturbed": {
            "pd": "推力增加轨迹偏移。\n撤力后逐步恢复。",
            "gc-pd": "受推时短暂偏离。\n本次撤力恢复最快。",
            "rbf": "受推时在线学补偿。\n偏移后仍能继续画。",
            "fuzzy": "按误差持续更新补偿。\n本次受推误差最小。",
        },
    }
    embedded = []
    for i, controller in enumerate(CONTROL):
        x, width = 0.82 + i * 2.97, 2.78
        video = ROOT / f"video/{controller}_{condition}.mp4"
        stamp = "113_9s" if normal else "55_5s"
        poster = ROOT / f"output/video_checks/{controller}_{condition}/frame_{stamp}.png"
        if not video.is_file() or not poster.is_file():
            raise FileNotFoundError(f"Missing original video or poster: {video}, {poster}")
        print(f"Embedding complete video: {video.name}", flush=True)
        movie = slide.shapes.add_movie(str(video), Inches(x), Inches(2.13),
                                       Inches(width), Inches(width * 0.75),
                                       poster_frame_image=str(poster), mime_type="video/mp4")
        movie.name = f"Embedded video / {controller} / {condition}"
        text(slide, x, 4.48, width, 0.48, LABEL[controller], 22, bold=True,
             align=PP_ALIGN.CENTER, name=f"Video title / {controller} / {condition}")
        text(slide, x, 5.09, width, 1.05, explanations[condition][controller], 18, MUTED,
             align=PP_ALIGN.CENTER, name=f"Video explanation / {controller} / {condition}")
        source_paths.extend((video, poster))
        embedded.append({"controller": controller, "condition": condition, "slide": number,
                         "source": str(video.relative_to(ROOT)),
                         "poster": str(poster.relative_to(ROOT)),
                         "explanation": explanations[condition][controller]})
    bottom = ("纸面线条显示实际 XY 投影；误差按真实三维笔尖位置计算。" if normal else
              "三次推力各持续 2 秒：20–22 秒、55–57 秒、90–92 秒。")
    text(slide, 0.82, 6.56, 11.7, 0.35, bottom, 17, MUTED)
    configure_parallel_video_autoplay(slide._element)
    return embedded


def build(destination):
    OUT.mkdir(parents=True, exist_ok=True)
    summary_path = ROOT / "output/summary.csv"
    manifest_path = ROOT / "output/experiment_manifest.json"
    manifest = json.loads(manifest_path.read_text())
    assert manifest["status"] == "completed", "PPT requires the completed eight-run batch"
    with summary_path.open(newline="") as stream:
        summary = {(r["controller"], r["condition"]): r for r in csv.DictReader(stream)}
    assert len(summary) == 8
    rmse = lambda c, condition: float(summary[c, condition]["drawing_tip_rmse_m"]) * 1000
    normal = {}
    source_paths = [summary_path, manifest_path]
    for controller in CONTROL:
        path = ROOT / f"output/normal/{controller}/trajectory.csv"
        print(f"Reading normal trajectory: {controller}", flush=True)
        normal[controller] = read_strokes(path)
        source_paths.append(path)
    forced_path = ROOT / "output/disturbed/fuzzy/trajectory.csv"
    print("Reading fuzzy disturbed trajectory", flush=True)
    forced = read_strokes(forced_path)
    source_paths.append(forced_path)
    deviations, recovery = {}, {}
    for c in CONTROL:
        csv_path = ROOT / f"output/comparisons/{c}/comparison.csv"
        json_path = csv_path.with_name("comparison_metrics.json")
        deviations[c] = read_deviation(csv_path)
        report = json.loads(json_path.read_text())
        recovery[c] = [window["recovery_time_s"] for window in report["recovery"]["windows"]]
        assert report["recovery"]["threshold_m"] == 0.002
        assert report["recovery"]["required_hold_s"] == 0.5
        source_paths.extend((csv_path, json_path))
    improvement = (1 - rmse("fuzzy", "disturbed") / rmse("gc-pd", "disturbed")) * 100

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "机械臂画家 · robot arm artist · raa-v0"
    prs.core_properties.subject = "Isaac Sim 七轴机械臂绘图与四种关节力矩控制的对比"
    prs.core_properties.author = ""
    prs.core_properties.keywords = "raa-v0, Isaac Sim, Panda, PD, RBF, 模糊自适应, 可编辑"
    prs.core_properties.comments = "Based on the completed eight-run experiment batch; all report metrics are pen-down 3D RMSE."
    scene = ROOT / "output/video_checks/fuzzy_normal/frame_113_9s.png"
    push_scene = ROOT / "output/video_checks/gc-pd_disturbed/frame_55_5s.png"

    # 1 — One clear purpose and an actual simulator image.
    s = base(prs, 1)
    text(s, 0.82, 1.55, 5.6, 0.98, "机械臂画家", 46, bold=True)
    text(s, 0.82, 2.63, 5.5, 0.57, "让七轴机械臂画出 ChatGPT 标志", 24)
    text(s, 0.82, 3.49, 5.35, 1.02, "画得准不准？\n被推一下，还能跟上吗？", 26, TEAL)
    text(s, 0.82, 5.73, 5.4, 0.75, "Isaac Sim 三维仿真\nrobot arm artist · raa-v0", 19, MUTED)
    picture(s, scene, 6.65, 1.66, 6.0, 4.5)

    # 2 — Native vertical flow and editable outline.
    s = base(prs, 2, "机械臂怎样画出图案？", "先安排笔尖怎么走，再算出每个关节怎么动。")
    steps = [("1", "读取 SVG 轮廓", "把标志变成一条条路线"),
             ("2", "安排笔尖运动", "平滑起停，笔画之间抬笔"),
             ("3", "计算关节角度", "连续逆运动学：位置 → 七个关节"),
             ("4", "控制关节力矩", "读回实际状态，及时纠正偏差"),
             ("5", "记录仿真结果", "保存真实轨迹、误差和视频")]
    for i, (number, title, detail) in enumerate(steps):
        y = 1.83 + i * 0.93
        rect(s, 0.82, y, 5.3, 0.75, WHITE, GRID)
        text(s, 1.00, y + 0.12, 0.35, 0.4, number, 23, TEAL, bold=True)
        text(s, 1.56, y + 0.04, 4.2, 0.36, title, 21, bold=True)
        text(s, 1.56, y + 0.42, 4.3, 0.28, detail, 16, MUTED)
        if i < 4:
            arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.34), Inches(y + 0.77),
                                       Inches(0.25), Inches(0.13))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = rgb(TEAL)
            arrow.line.fill.background()
    outline_plot(s, 7.23, 1.73, 4.85, 4.35, "目标轮廓", [("目标", normal["fuzzy"]["target"], TEAL, None)], axes=False)
    text(s, 7.23, 6.2, 4.85, 0.45, "Franka Panda 七轴  ·  16 cm  ·  8 笔", 18, MUTED,
         align=PP_ALIGN.CENTER)

    # 3 — Plain mechanisms first, one simple professional expression second.
    s = base(prs, 3, "四种控制器，区别在哪里？", "都在纠正偏差，区别是怎样补上额外的力矩。")
    rows = [["控制器", "通俗理解"],
            ["PD", "偏了就纠正，动得太急就缓一下"],
            ["重力补偿 PD", "纠正偏差时，先托住机械臂的重量"],
            ["RBF 自适应", "根据关节位置和速度，边画边学补偿"],
            ["模糊自适应", "把偏差和变化速度分档，用规则学补偿"]]
    shape = table(s, 0.82, 1.94, 11.7, 3.40, rows, [0.25, 0.75], 22)
    # Explanations are easier to read left-aligned.
    for row in list(shape.table.rows)[1:]:
        for p in row.cells[1].text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
    text(s, 0.82, 5.70, 11.7, 0.51, "RBF 与模糊：重力补偿 PD ＋ 在线学习补偿", 24, TEAL, bold=True)
    text(s, 0.82, 6.30, 11.7, 0.42, "每次从零权重开始；四种方案都输出七个关节的力矩。", 19, MUTED)

    # 4 — Fair test and exact physical force vectors.
    s = base(prs, 4, "同一张图，同样的三次推力", "相同路线、起点和 PD 参数；4 种方法 × 正常 / 受推，共 8 组。")
    picture(s, push_scene, 0.82, 1.89, 5.80, 4.35)
    text(s, 0.82, 6.40, 5.80, 0.32, "真实笔尖处施力，目标路线保持原样", 18, MUTED, align=PP_ALIGN.CENTER)
    text(s, 7.06, 1.87, 5.3, 0.42, "三次推力，每次持续 2 秒", 23, bold=True)
    table(s, 7.06, 2.47, 5.43, 2.59,
          [["时段（s）", "方向", "力（N）"],
           ["20–22", "+X", "(3, 0, 0)"],
           ["55–57", "−Y", "(0, −3, 0)"],
           ["90–92", "+X、+Y", "(2, 2, 0)"]], [0.31, 0.26, 0.43], 19)
    text(s, 7.06, 5.43, 5.4, 0.90, "每 1 ms 更新一次控制\n每组完整运动约 114 秒", 21, MUTED)

    # 5 — Four same-scale outline plots, exact metrics in an editable table.
    s = base(prs, 5, "正常绘画：补偿重力后更准确")
    legend(s, 0.84, 1.35, [("目标", TARGET, MSO_LINE_DASH_STYLE.DASH), ("实际", TEAL, None)], 16)
    for i, c in enumerate(CONTROL):
        outline_plot(s, 0.62 + i * 3.10, 1.89, 2.80, 3.36, SHORT[c],
                     [("目标", normal[c]["target"], TARGET, MSO_LINE_DASH_STYLE.DASH),
                      ("实际", normal[c]["actual"], TEAL, None)], small=True)
    text(s, 0.82, 5.48, 11.7, 0.42, "落笔阶段三维误差 RMSE（mm，越小越好）", 20)
    table(s, 0.82, 6.00, 11.7, 0.74,
          [[f"{SHORT[c]}\n{rmse(c, 'normal'):.3f}" for c in CONTROL]], [0.25] * 4, 20)

    # 6 — Pairing table and editable normal / disturbed / target XY plot.
    s = base(prs, 6, "受到推力：模糊的整体误差更小", "目标、无干扰轨迹和受扰轨迹，按同一时刻比较。")
    text(s, 0.82, 1.85, 5.5, 0.42, "落笔阶段三维 RMSE（mm）", 21, bold=True)
    table(s, 0.82, 2.43, 5.57, 2.82,
          [["控制器", "正常", "受推"]] + [[SHORT[c], f"{rmse(c, 'normal'):.3f}",
                                            f"{rmse(c, 'disturbed'):.3f}"] for c in CONTROL],
          [0.44, 0.28, 0.28], 20, highlight_row=4)
    text(s, 0.82, 5.64, 5.56, 1.02,
         f"模糊受推 RMSE 为 1.77 mm\n比重力补偿 PD 低约 {improvement:.0f}%", 23, TEAL, bold=True)
    outline_plot(s, 7.0, 1.74, 5.45, 4.48, "模糊控制的轨迹对比",
                 [("目标", normal["fuzzy"]["target"], TARGET, MSO_LINE_DASH_STYLE.DASH),
                  ("正常", normal["fuzzy"]["actual"], BLUE, None),
                  ("受推", forced["actual"], ORANGE, None)])
    legend(s, 7.43, 6.53, [("目标", TARGET, MSO_LINE_DASH_STYLE.DASH),
                           ("正常", BLUE, None), ("受推", ORANGE, None)], 15, 1.55)

    # 7 — True editable chart with embedded workbook + exact recovery table.
    s = base(prs, 7, "撤去推力：重力补偿 PD 恢复最快", "看受推轨迹多久回到原来的无干扰轨迹附近。")
    deviation_chart(s, 0.82, 2.45, 5.80, 3.68, deviations)
    text(s, 7.02, 1.98, 5.45, 0.45, "撤力后的恢复时间（s）", 21, bold=True)
    table(s, 7.02, 2.53, 5.46, 2.90,
          [["控制器", "第 1 次", "第 2 次", "第 3 次"]] +
          [[SHORT[c]] + [f"{value:.3f}" for value in recovery[c]] for c in CONTROL],
          [0.37, 0.21, 0.21, 0.21], 18, highlight_row=2)
    text(s, 7.02, 5.75, 5.46, 0.55, "重力补偿 PD：三次均约 0.2 秒", 21, TEAL, bold=True)
    text(s, 0.82, 6.51, 11.7, 0.40,
         "恢复标准：回到无干扰轨迹的 2 mm 内，并连续保持至少 0.5 秒。", 18, MUTED)

    # 8–9 — Native embedded MP4 objects, no external video file references.
    embedded_videos = video_slide(prs, 8, "normal", source_paths)
    embedded_videos += video_slide(prs, 9, "disturbed", source_paths)

    # 10 — Three supported takeaways and a single next step.
    s = base(prs, 10, "本次实验的三个结论")
    takeaways = [
        ("01", "先托住自重，画得更准", "正常三维 RMSE：PD 约 33 mm → 重力补偿 PD 约 0.56 mm"),
        ("02", "模糊控制的整体误差最小", "本次受推 RMSE 约 1.77 mm，比重力补偿 PD 低约 31%"),
        ("03", "重力补偿 PD 撤力恢复最快", "三次均约 0.2 秒；绘画精度和恢复速度要分别看"),
    ]
    for i, (number, title, explanation) in enumerate(takeaways):
        y = 1.81 + i * 1.34
        text(s, 0.82, y + 0.03, 0.8, 0.64, number, 32, TEAL, bold=True)
        text(s, 1.95, y, 10.55, 0.53, title, 27, bold=True)
        text(s, 1.95, y + 0.68, 10.55, 0.43, explanation, 20, MUTED)
        if i < 2:
            line(s, 1.95, y + 1.19, 12.48, y + 1.19, GRID, 0.7)
    text(s, 0.82, 6.10, 11.7, 0.37, "下一步：从图片提取轮廓，再调整参数与推力强度。", 20, TEAL)
    text(s, 0.82, 6.62, 11.7, 0.29, "结论对应本次图案、控制参数与外力设置。", 16, MUTED)

    notes = load_notes()
    fallback_notes = {
        1: "项目完全在 Isaac Sim 中完成，使用 Franka Panda 七轴机械臂描绘 ChatGPT 标志的八条轮廓。",
        2: "SVG 轮廓通过平滑时间规划、连续逆运动学得到参考关节角度，再用力矩控制执行。目标轮廓来源：output/normal/fuzzy/trajectory.csv。",
        3: "PD 力矩为 Kp(q_ref-q)+Kd(dq_ref-dq)。GC-PD 增加 g(q)，两种自适应方法再增加 W^T phi，并从零权重在线更新。",
        4: "外力为世界坐标系下的 (3,0,0)、(0,-3,0)、(2,2,0) N，分别持续20–22、55–57、90–92s。第三次合力约2.828N。",
        5: "表中 RMSE=sqrt(mean(||p_actual-p_target||^2))，只统计落笔时刻。轨迹按世界坐标XY显示，同轴同尺度。source: output/summary.csv。",
        6: "正常与受推记录逐时刻对齐。模糊的落笔三维 RMSE 为1.768879888mm，相对GC-PD的2.558163937mm降低30.8535%。PD受扰RMSE稍低，不能据此判为更抗扰。",
        7: "恢复依据受推位置减无干扰位置的三维范数，撤力后连续0.5s不超过2mm。恢复时间记录这段区间的开始时间减撤力时间。曲线19–27s，展示第一次推力。",
        8: "四个正常绘画视频完整内嵌在PPT中，进入本页后四个视频自动同时播放。PD轮廓偏移明显，补偿方案轮廓贴近目标。墨迹为实际XY投影，XYZ数据保持真实。模拟虚拟笔尖，没有模拟笔纸接触压力。",
        9: "四个受扰视频完整内嵌在PPT中。推力发生在20–22、55–57、90–92秒。GC-PD撤力恢复最快，模糊受扰整体RMSE最低。PD箭头被桌面遮挡，实际施力和轨迹偏移已经核验。",
        10: "重力补偿改善精度；本次参数下模糊整体RMSE最低；GC-PD撤力恢复最快。结论针对这一图案与推力参数。",
    }
    for i, slide in enumerate(prs.slides, 1):
        body = notes.get(i, fallback_notes[i])
        slide.notes_slide.notes_text_frame.text = body + "\n\n编辑说明：文字、表格、流程图、轨迹线及时间曲线均为原生 PowerPoint 对象。时间曲线可通过“编辑数据”修改；轨迹线可通过“编辑顶点”修改。真实仿真截图为图片。第8、9页的八个视频完整内嵌，视频下方说明为可编辑文本框。"
    assert len(prs.slides) == TOTAL_SLIDES
    # Disable template-inherited shadows, including native freeform paths.
    for slide in prs.slides:
        for shape in slide.shapes:
            for effect_ref in shape._element.xpath(".//a:effectRef"):
                effect_ref.set("idx", "0")
            for props in shape._element.xpath("./p:spPr"):
                if props.find("a:effectLst", props.nsmap) is None:
                    props.append(OxmlElement("a:effectLst"))
    # Simple static shape bounds check, followed by actual rendering outside here.
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= -1000 and shape.top >= -1000, (i, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 1000, (i, shape.name, "right edge")
            assert shape.top + shape.height <= prs.slide_height + 1000, (i, shape.name, "bottom edge")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.building.pptx")
    try:
        prs.save(temporary)
        temporary.replace(destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    source_hashes = {}
    for p in source_paths + [scene, push_scene]:
        digest = hashlib.sha256()
        with p.open("rb") as stream:
            while block := stream.read(1024 * 1024):
                digest.update(block)
        source_hashes[str(p.relative_to(ROOT))] = digest.hexdigest()
    audit = {
        "presentation": str(destination.relative_to(ROOT)) if destination.is_relative_to(ROOT) else str(destination),
        "slide_count": TOTAL_SLIDES,
        "slide_size": "16:9",
        "font": FONT,
        "metric": "pen-down 3D position RMSE, mm",
        "normal_rmse_mm": {c: rmse(c, "normal") for c in CONTROL},
        "disturbed_rmse_mm": {c: rmse(c, "disturbed") for c in CONTROL},
        "recovery_s": recovery,
        "fuzzy_vs_gc_pd_disturbed_rmse_reduction_percent": improvement,
        "recovery_definition": "After force off, disturbed-normal 3D distance <= 2 mm for at least 0.5 s; time is start of qualifying interval minus force end.",
        "outline_display": "10 ms displayed samples, stroke endpoints preserved, RDP tolerance 0.05 mm; all reported RMSE use complete 1 ms experiment data",
        "time_series_display": "first force window, 19–27 s, displayed samples every 10 ms",
        "native_editable": ["text", "tables", "flowchart", "trajectory freeform paths", "XY time-series chart with embedded Excel workbook"],
        "raster_assets": [str(p.relative_to(ROOT)) for p in (scene, push_scene)],
        "embedded_videos": embedded_videos,
        "video_delivery": "8 complete internally embedded MP4 files; four videos start automatically in parallel on each video slide; no external video links",
        "video_playback": {"trigger": "slide entry", "start_delay_ms": 0,
                           "mode": "parallel", "videos_per_slide": 4, "slides": [8, 9]},
        "pdf_preview": "Static poster frames only; embedded videos play in the PPTX slideshow",
        "source_sha256": source_hashes,
        "notes_from_markdown": sorted(notes),
        "shape_counts": [len(s.shapes) for s in prs.slides],
    }
    (OUT / "presentation_sources.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    print(f"Saved editable presentation: {destination}", flush=True)
    print(f"Ten slides; eight fully embedded videos; native chart workbook; notes for {len(notes)}/{TOTAL_SLIDES} slides.", flush=True)
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUT / "raa-v0.pptx")
    parser.add_argument("--sync-videos", action="store_true",
                        help="Update only the existing PPTX: four videos automatically start together on each video slide")
    args = parser.parse_args()
    destination = args.output.expanduser().resolve()
    if args.sync_videos:
        synchronize_existing_presentation(destination)
    else:
        build(destination)


if __name__ == "__main__":
    main()
